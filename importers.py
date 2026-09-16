"""POS ve hakediş dosyalarını okur: CSV / XLSX.

Türkçe dosyaların klasik tuzakları burada tek yerde çözülür:
  - ondalık ayracı virgül  ("1.234,56" -> 1234.56)
  - tarih formatı gg.aa.yyyy
  - CSV ayracı noktalı virgül
  - cp1254 / latin-5 kodlama
POS'tan POS'a başlıklar değiştiği için sütun eşlemesi kullanıcıdan alınır;
otomatik tahmin sadece ön dolgu yapar.
"""
import csv
import io
import unicodedata
import os
import re
from datetime import datetime

try:
    import openpyxl
    HAVE_XLSX = True
except ImportError:
    HAVE_XLSX = False


# --------------------------------------------------------------------
# Değer dönüştürücüler
# --------------------------------------------------------------------
_NUM_CLEAN = re.compile(r"[^\d,.\-]")


def to_float(v):
    """'1.234,56 TL' -> 1234.56 ; '' -> 0.0

    Hem TR (1.234,56) hem EN (1,234.56) biçimini ayırt eder: son görülen
    ayraç ondalık kabul edilir.
    """
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = _NUM_CLEAN.sub("", str(v).strip())
    if not s or s in ("-", ".", ","):
        return 0.0
    last_c, last_d = s.rfind(","), s.rfind(".")
    if last_c > last_d:            # TR: virgül ondalık
        s = s.replace(".", "").replace(",", ".")
    elif last_d > last_c:          # EN: nokta ondalık
        s = s.replace(",", "")
    else:                          # tek tip ayraç yok
        s = s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


_DATE_FORMATS = ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y",
                 "%Y/%m/%d", "%Y.%m.%d", "%d.%m.%y")


def to_date(v):
    """Çeşitli formatları ISO 'YYYY-MM-DD'ye çevirir. Saat kısmı atılır."""
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if hasattr(v, "isoformat") and not isinstance(v, str):
        return v.isoformat()[:10]
    s = str(v).strip()
    s = re.split(r"[ T]", s)[0]
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return s[:10] if len(s) >= 10 else s


def norm_status(v):
    """Durum metnini standart koda indirger.

    Karşılaştırma _norm() üzerinden yapılır: "İptal Edildi".lower() sonucu
    'i' + birleşen nokta içerdiği için düz `in` kontrolü tutmaz ve iptal
    siparişler sessizce 'teslim edildi' sayılır — mutabakatı baştan bozar.
    """
    s = _norm(v)
    if not s:
        return "delivered"
    if any(k in s for k in ("iptal", "cancel", "reddedil", "rejected", "basarisiz")):
        return "cancelled"
    if any(k in s for k in ("iade", "refund", "geri odendi", "geri ode")):
        return "refunded"
    return "delivered"


def norm_payment(v):
    """Ödeme yöntemini standart koda indirger.

    Kapıda ödeme ile online ödemeyi karıştırmak mutabakatı tamamen bozar:
    kapıda ödemede parayı biz aldık, platform bize ciro ödemez.
    """
    s = _norm(v)
    if not s:
        return "online"
    if any(k in s for k in ("nakit", "cash")):
        return "cash"
    if any(k in s for k in ("kapida", "on delivery", "kapi")):
        return "card_on_delivery"
    return "online"


# --------------------------------------------------------------------
# Dosya okuma
# --------------------------------------------------------------------
def read_table(path_or_bytes, filename=""):
    """Dosyayı (başlıklar, satırlar) olarak döndürür. Satırlar dict."""
    name = (filename or str(path_or_bytes)).lower()
    if isinstance(path_or_bytes, bytes):
        data = path_or_bytes
    else:
        with open(path_or_bytes, "rb") as f:
            data = f.read()

    if name.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(data)
    return _read_csv(data)


def _read_xlsx(data):
    if not HAVE_XLSX:
        raise RuntimeError("Excel okumak için openpyxl gerekli: pip install openpyxl")
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    headers = None
    out = []
    for raw in rows_iter:
        if headers is None:
            # İlk dolu satırı başlık kabul et (üstte boş/başlık süsü olabilir)
            if raw and any(c not in (None, "") for c in raw):
                headers = [str(c).strip() if c is not None else "" for c in raw]
            continue
        if raw is None or all(c in (None, "") for c in raw):
            continue
        out.append({headers[i]: raw[i] for i in range(min(len(headers), len(raw)))})
    wb.close()
    return headers or [], out


def _read_csv(data):
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1254", "iso-8859-9", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = data.decode("utf-8", errors="replace")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ";" if sample.count(";") > sample.count(",") else ","

    reader = csv.DictReader(io.StringIO(text), delimiter=delim)
    headers = [h.strip() for h in (reader.fieldnames or [])]
    rows = []
    for r in reader:
        rows.append({(k.strip() if k else ""): v for k, v in r.items()})
    return headers, rows


# --------------------------------------------------------------------
# Otomatik sütun tahmini
# --------------------------------------------------------------------
POS_SYNONYMS = {
    "platform_order_id": ["sipariş no", "siparis no", "sipariş numarası", "order id",
                          "order_id", "platform sipariş", "harici no", "external id",
                          "sipariş kodu", "order number"],
    "order_date": ["tarih", "sipariş tarihi", "siparis tarihi", "order date",
                   "date", "işlem tarihi", "created"],
    "items_subtotal": ["ürün tutarı", "urun tutari", "ara toplam", "sepet tutarı",
                       "subtotal", "tutar", "toplam", "items", "ürün toplamı"],
    "delivery_fee": ["teslimat", "teslimat ücreti", "kurye", "delivery fee",
                     "gönderim", "servis ücreti"],
    "discount_restaurant": ["restoran indirimi", "indirim", "restoran kampanya",
                            "restaurant discount", "bizim indirim"],
    "discount_platform": ["platform indirimi", "platform kampanya",
                          "platform discount", "kupon"],
    "tip": ["bahşiş", "bahsis", "tip"],
    "payment_method": ["ödeme", "ödeme tipi", "ödeme yöntemi", "payment",
                       "payment method", "odeme"],
    "status": ["durum", "sipariş durumu", "status", "state"],
}

PAYOUT_SYNONYMS = {
    "platform_order_id": ["sipariş no", "siparis no", "order id", "order_id",
                          "sipariş numarası", "sipariş kodu", "order number"],
    "order_date": ["tarih", "sipariş tarihi", "order date", "date"],
    "gross_reported": ["brüt", "brut", "brüt tutar", "sipariş tutarı", "gross",
                       "toplam tutar", "ciro"],
    "commission_charged": ["komisyon", "komisyon tutarı", "commission",
                           "hizmet bedeli", "komisyon bedeli"],
    "commission_vat_charged": ["komisyon kdv", "kdv", "vat", "komisyon kdv tutarı"],
    "other_deductions": ["kesinti", "diğer kesinti", "deduction", "mahsup",
                         "diger kesinti"],
    "net_paid": ["net", "net tutar", "ödenen", "odenen", "hakediş", "hakedis",
                 "net ödeme", "payout", "net_paid"],
    "line_note": ["açıklama", "aciklama", "not", "note", "description"],
}


def _norm(s):
    """Başlığı karşılaştırılabilir hale getirir: Türkçe harfler sadeleşir.

    Dikkat: Python'da "İ".lower() sonucu 'i' + ayrı bir birleşen nokta
    karakteridir (U+0307). Bu karakter temizlenince kelime ikiye bölünür
    ("Restoran İndirimi" -> "restoran i ndirimi") ve eşleşme kaçar.
    Bu yüzden Türkçe harfler küçültmeden ÖNCE sadeleştirilir ve kalan
    birleşen işaretler ayrıca atılır.
    """
    s = str(s or "").strip()
    for a, b in (("İ", "I"), ("I", "I"), ("ı", "i"), ("Ş", "S"), ("ş", "s"),
                 ("Ğ", "G"), ("ğ", "g"), ("Ü", "U"), ("ü", "u"),
                 ("Ö", "O"), ("ö", "o"), ("Ç", "C"), ("ç", "c")):
        s = s.replace(a, b)
    s = s.lower()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def _score(nh, nc):
    """Başlık ile aday terimin benzerlik puanı. 0 = eşleşme yok.

    Kelime sınırına dikkat edilir: aksi halde 'tip' (bahşiş) terimi
    'Ödeme Tipi' sütununu kapar ve bahşiş sütunu ödeme yöntemi sanılır.
    """
    if not nh or not nc:
        return 0
    if nh == nc:
        return 100
    # Aday, başlıkta tam kelime olarak geçiyor mu
    if re.search(r"\b%s\b" % re.escape(nc), nh):
        return 80
    # Türkçe ek toleransı: 'indirim' -> 'indirimi', 'tutar' -> 'tutari'
    for w in nh.split():
        if len(nc) >= 4 and w.startswith(nc) and len(w) - len(nc) <= 3:
            return 65
        if len(w) >= 4 and nc.startswith(w) and len(nc) - len(w) <= 3:
            return 60
    return 0


def guess_mapping(headers, kind="pos_orders"):
    """Başlıklara bakıp hedef alan -> sütun tahmini üretir.

    Tüm (alan, sütun) çiftleri puanlanıp en yüksek puandan başlayarak
    atanır. Alan alan ilerlemek yanlış sonuç verir: önce işlenen zayıf
    bir eşleşme, sonraki alanın kesin eşleşmesini elinden alır.

    Tahmin sadece ön dolgudur; kullanıcı panelde onaylar veya düzeltir.
    Yanlış eşlenen sütun sessizce sıfır okunur ve mutabakatı bozar,
    o yüzden son karar kullanıcıda kalmalı.
    """
    syn = POS_SYNONYMS if kind == "pos_orders" else PAYOUT_SYNONYMS
    nheaders = {h: _norm(h) for h in headers if h}

    pairs = []
    for field, cands in syn.items():
        for h, nh in nheaders.items():
            best = max((_score(nh, _norm(c)) for c in cands), default=0)
            if best >= 60:
                pairs.append((best, field, h))

    pairs.sort(key=lambda p: (-p[0], p[1]))
    mapping, used_fields, used_cols = {}, set(), set()
    for score, field, h in pairs:
        if field in used_fields or h in used_cols:
            continue
        mapping[field] = h
        used_fields.add(field)
        used_cols.add(h)
    return mapping


# --------------------------------------------------------------------
# Eşlemeye göre kayda dönüştürme
# --------------------------------------------------------------------
def build_orders(rows, mapping, platform, source_file=""):
    """Ham satırları orders tablosu kayıtlarına çevirir."""
    out, skipped = [], []
    for i, r in enumerate(rows, start=2):
        oid = str(r.get(mapping.get("platform_order_id", ""), "") or "").strip()
        if not oid:
            skipped.append((i, "sipariş no boş"))
            continue
        d = to_date(r.get(mapping.get("order_date", ""), ""))
        if not d:
            skipped.append((i, "tarih okunamadı"))
            continue
        out.append({
            "platform": platform,
            "platform_order_id": oid,
            "order_date": d,
            "items_subtotal": to_float(r.get(mapping.get("items_subtotal", ""), 0)),
            "delivery_fee": to_float(r.get(mapping.get("delivery_fee", ""), 0)),
            "discount_restaurant": to_float(r.get(mapping.get("discount_restaurant", ""), 0)),
            "discount_platform": to_float(r.get(mapping.get("discount_platform", ""), 0)),
            "tip": to_float(r.get(mapping.get("tip", ""), 0)),
            "payment_method": norm_payment(r.get(mapping.get("payment_method", ""), "")),
            "status": norm_status(r.get(mapping.get("status", ""), "")),
            "source_file": source_file,
        })
    return out, skipped


def build_payout_lines(rows, mapping, platform):
    """Ham satırları payout_lines kayıtlarına çevirir."""
    out, skipped = [], []
    for i, r in enumerate(rows, start=2):
        oid = str(r.get(mapping.get("platform_order_id", ""), "") or "").strip()
        if not oid:
            skipped.append((i, "sipariş no boş"))
            continue
        com = abs(to_float(r.get(mapping.get("commission_charged", ""), 0)))
        out.append({
            "platform": platform,
            "platform_order_id": oid,
            "order_date": to_date(r.get(mapping.get("order_date", ""), "")),
            "gross_reported": to_float(r.get(mapping.get("gross_reported", ""), 0)),
            "commission_charged": com,
            "commission_vat_charged": abs(to_float(
                r.get(mapping.get("commission_vat_charged", ""), 0))),
            "other_deductions": abs(to_float(
                r.get(mapping.get("other_deductions", ""), 0))),
            "net_paid": to_float(r.get(mapping.get("net_paid", ""), 0)),
            "line_note": str(r.get(mapping.get("line_note", ""), "") or "").strip(),
        })
    return out, skipped
